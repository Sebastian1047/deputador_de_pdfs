"""Excel Online en OneDrive/SharePoint mediante Graph + openpyxl, sin disco.

Usa permisos de aplicación para operaciones sobre el archivo completo. No llama
a endpoints /workbook (esos requieren permisos delegados). Para libros sencillos
.xlsx con una hoja de registros dedicada; no utilizar en libros con macros ni con
funciones avanzadas sin validar primero la integridad en una copia de prueba.
"""
from __future__ import annotations

import io
import os
import re
from typing import Sequence
from urllib.parse import quote, urlsplit

from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

from .core import Campo

GRAPH = "https://graph.microsoft.com/v1.0"
MAX_XLSX_BYTES = 20 * 1024 * 1024
MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _id_valido(value: str, nombre: str) -> str:
    if not value or len(value) > 300 or not re.fullmatch(r"[A-Za-z0-9!_\-.:{}]+", value):
        raise ValueError(f"Identificador {nombre} no válido")
    return quote(value, safe="")


class GraphExcelFile:
    """Peticiones al archivo específico. El token solo se envía a graph.microsoft.com."""

    def __init__(self, graph_client, drive_id: str, item_id: str):
        self.auth = graph_client.auth
        self.http = graph_client.http
        drive = _id_valido(drive_id, "drive")
        item = _id_valido(item_id, "item")
        self.base = f"{GRAPH}/drives/{drive}/items/{item}"

    def _token(self) -> str:
        resultado = self.auth.acquire_token_for_client(
            scopes=["https://graph.microsoft.com/.default"]
        )
        token = resultado.get("access_token")
        if not token:
            raise RuntimeError("Microsoft 365 no autorizó acceso al Excel")
        return token

    def _peticion_graph(self, method, url: str, **kwargs):
        parsed = urlsplit(url)
        if parsed.scheme != "https" or parsed.netloc != "graph.microsoft.com":
            raise ValueError("Destino Graph inválido")
        headers = {"Authorization": "Bearer " + self._token(), **kwargs.pop("headers", {})}
        request_fn = getattr(self.http, method)
        respuesta = request_fn(
            url, headers=headers, timeout=(10, 90), allow_redirects=False, **kwargs
        )
        return respuesta

    def _metadata(self) -> dict:
        r = self._peticion_graph("get", self.base)
        if r.status_code != 200:
            raise RuntimeError(f"No pudo acceder al archivo Excel (HTTP {r.status_code})")
        metadata = r.json()
        if not str(metadata.get("name", "")).lower().endswith(".xlsx"):
            raise ValueError("El destino debe ser un libro Excel .xlsx")
        if "file" not in metadata:
            raise ValueError("El elemento seleccionado no es un archivo")
        if int(metadata.get("size") or 0) > MAX_XLSX_BYTES:
            raise ValueError("Excel supera el tamaño permitido (20 MiB)")
        if not metadata.get("eTag"):
            raise ValueError("Excel no suministró su versión eTag")
        return metadata

    def leer(self) -> tuple[bytes, str]:
        """Lee una versión en RAM, verificando tamaño y dominio de descarga."""
        metadata = self._metadata()
        origen = metadata.get("@microsoft.graph.downloadUrl", "")
        parsed = urlsplit(origen)
        hostname = (parsed.hostname or "").lower()
        if (
            parsed.scheme != "https"
            or not (
                hostname.endswith(".sharepoint.com")
                or hostname.endswith(".sharepointonline.com")
                or hostname.endswith(".1drv.com")
            )
            or parsed.username
            or parsed.password
        ):
            raise ValueError("Microsoft no devolvió una URL de descarga segura")
        # IMPORTANTE: la URL ya está autorizada por Microsoft; jamás pasar bearer.
        r = self.http.get(
            origen, timeout=(10, 90), allow_redirects=False,
            headers={"Accept": MIME},
        )
        if r.status_code != 200:
            raise RuntimeError(f"No pudo descargar el Excel (HTTP {r.status_code})")
        if len(r.content) > MAX_XLSX_BYTES or not r.content.startswith(b"PK"):
            raise ValueError("El archivo descargado no es un .xlsx válido o es demasiado grande")
        return r.content, metadata["eTag"]

    def guardar(self, contenido: bytes, etag: str) -> None:
        if len(contenido) > MAX_XLSX_BYTES or not contenido.startswith(b"PK"):
            raise ValueError("Archivo de salida demasiado grande o inválido")
        # Solo actualiza si nadie editó el Excel desde la descarga.
        actual = self._metadata()
        if actual["eTag"] != etag:
            raise RuntimeError("Conflicto: Excel fue modificado por otra persona; vuelva a procesar")
        respuesta = self._peticion_graph(
            "put", self.base + "/content",
            headers={"If-Match": etag, "Content-Type": MIME},
            data=contenido,
        )
        if respuesta.status_code == 412:
            raise RuntimeError("Conflicto: Excel cambió durante el guardado; reintente")
        if respuesta.status_code not in (200, 201):
            raise RuntimeError(f"No pudo guardar el Excel (HTTP {respuesta.status_code})")


class ExcelOnlineWriter:
    """Acumula hasta 40+ registros en memoria y guarda el Excel una sola vez."""

    def __init__(self, archivo: GraphExcelFile, hoja: str = "Registros"):
        self.archivo = archivo
        self.hoja = hoja
        self.wb = None
        self.ws = None
        self.etag = None
        self.agregados = 0
        self.tabla = None
        self.ancho = None

    def claves_existentes(self, campos: Sequence[Campo]) -> set[str]:
        datos, self.etag = self.archivo.leer()
        self.wb = load_workbook(io.BytesIO(datos), read_only=False, data_only=False, keep_links=True)
        if self.wb.vba_archive is not None:
            raise ValueError("No se admiten libros con macros")
        if self.hoja not in self.wb.sheetnames:
            raise ValueError(f"No existe la hoja requerida: {self.hoja}")
        self.ws = self.wb[self.hoja]
        columnas = ["document_key", "processed_utc", *[campo.nombre for campo in campos]]
        self.ancho = len(columnas)
        if len(columnas) > 16384:
            raise ValueError("Demasiadas columnas")
        actuales = [self.ws.cell(1, i + 1).value for i in range(len(columnas))]
        if all(valor is None for valor in actuales) and self.ws.max_row <= 1:
            for i, columna in enumerate(columnas, start=1):
                self.ws.cell(1, i, columna)
        elif actuales != columnas:
            raise ValueError("Las columnas del Excel no coinciden con las reglas de extracción")
        # Nunca insertar silenciosamente en una hoja con otras columnas o tablas.
        extra = self.ws.cell(1, len(columnas) + 1).value
        if extra is not None:
            raise ValueError("La hoja tiene columnas adicionales no esperadas")
        tablas = list(self.ws.tables.values())
        if len(tablas) > 1:
            raise ValueError("Se necesita una sola tabla en la hoja de destino")
        if tablas:
            tabla = tablas[0]
            if not tabla.ref.startswith("A1:"):
                raise ValueError("La tabla debe comenzar en A1")
            self.tabla = tabla
        else:
            self.tabla = None
        return {
            str(self.ws.cell(row, 1).value)
            for row in range(2, self.ws.max_row + 1)
            if self.ws.cell(row, 1).value not in (None, "")
        }

    def agregar(self, clave: str, fecha_utc: str, datos: Sequence[str]) -> None:
        if self.ws is None:
            raise RuntimeError("Se debe leer el libro antes de insertar filas")
        if len(datos) != self.ancho - 2:
            raise ValueError("Cantidad de valores distinta a las columnas del Excel")
        def texto(valor: object) -> str:
            valor = str(valor)
            # No evaluar fórmulas derivadas de PDF o mensajes externos.
            return "'" + valor if valor.startswith(("=", "+", "-", "@")) else valor
        self.ws.append([texto(clave), fecha_utc, *[texto(v) for v in datos]])
        self.agregados += 1

    def finalizar(self) -> None:
        if self.ws is None or self.wb is None or self.etag is None:
            raise RuntimeError("El libro no está cargado")
        if not self.agregados:
            return
        from openpyxl.utils import get_column_letter
        final = f"A1:{get_column_letter(self.ancho)}{self.ws.max_row}"
        if self.tabla:
            self.tabla.ref = final
            if self.tabla.autoFilter:
                self.tabla.autoFilter.ref = final
        else:
            tabla = Table(displayName="RegistrosPDF", ref=final)
            tabla.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2", showFirstColumn=False,
                showLastColumn=False, showRowStripes=True, showColumnStripes=False,
            )
            self.ws.add_table(tabla)
        salida = io.BytesIO()
        self.wb.save(salida)
        self.archivo.guardar(salida.getvalue(), self.etag)


def crear_excel_online(graph_client):
    requeridas = ("MS_EXCEL_DRIVE_ID", "MS_EXCEL_ITEM_ID")
    faltan = [clave for clave in requeridas if not os.getenv(clave)]
    if faltan:
        raise RuntimeError("Faltan variables del Excel Online: " + ", ".join(faltan))
    return ExcelOnlineWriter(
        GraphExcelFile(
            graph_client,
            os.environ["MS_EXCEL_DRIVE_ID"],
            os.environ["MS_EXCEL_ITEM_ID"],
        ),
        os.getenv("MS_EXCEL_SHEET") or "Registros",
    )


def resolver_ubicacion_onedrive(graph_client, propietario: str, ruta: str) -> tuple[str, str]:
    """Obtiene driveId e itemId usando propietario y ruta relativos a su OneDrive."""
    propietario = propietario.strip()
    if (
        not propietario or len(propietario) > 254 or "@" not in propietario
        or any(c in propietario for c in ("/", "\\", "?", "#", " ", "%"))
    ):
        raise ValueError("MS_ONEDRIVE_OWNER debe ser una cuenta empresarial válida")
    ruta = ruta.strip().strip("/")
    partes = ruta.split("/")
    if (
        not ruta or len(ruta) > 700 or not ruta.lower().endswith(".xlsx")
        or any(p in ("", ".", "..") or "\\" in p or "\x00" in p for p in partes)
    ):
        raise ValueError("MS_ONEDRIVE_PATH debe apuntar a un archivo .xlsx válido")
    propietario_encoded = quote(propietario, safe="")
    archivo_encoded = "/".join(quote(parte, safe="") for parte in partes)
    url = f"{GRAPH}/users/{propietario_encoded}/drive/root:/{archivo_encoded}:"
    meta = graph_client.get(url)
    if not isinstance(meta, dict) or "file" not in meta:
        raise ValueError("La ubicación en OneDrive no corresponde a un archivo")
    if not str(meta.get("name", "")).lower().endswith(".xlsx"):
        raise ValueError("El archivo de destino debe ser .xlsx")
    drive_id = (meta.get("parentReference") or {}).get("driveId")
    item_id = meta.get("id")
    _id_valido(str(drive_id or ""), "drive")
    _id_valido(str(item_id or ""), "item")
    return str(drive_id), str(item_id)


def crear_excel_desde_onedrive(graph_client, propietario: str, ruta: str,
                              hoja: str = "Registros"):
    drive_id, item_id = resolver_ubicacion_onedrive(graph_client, propietario, ruta)
    return ExcelOnlineWriter(GraphExcelFile(graph_client, drive_id, item_id), hoja)
