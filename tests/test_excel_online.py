"""Pruebas con un Excel xlsx en memoria, sin permisos ni redes reales."""
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from depurador.core import Campo, procesar, AdjuntoPDF
from depurador.excel_online import ExcelOnlineWriter, GraphExcelFile, MAX_XLSX_BYTES


def ejemplo_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = "Registros"
    ws.append(["document_key", "processed_utc", "numero"])
    ws.append(["outlook:demo:mensaje-1:adjunto-1", "2026-10-20", "101"])
    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()


class ArchivoSimulado:
    def __init__(self):
        self.datos = ejemplo_xlsx()
        self.etag = '"1"'
        self.guardados = []

    def leer(self):
        return self.datos, self.etag

    def guardar(self, contenido, etag):
        assert etag == self.etag
        self.guardados.append(contenido)
        self.datos = contenido
        self.etag = '"2"'


def test_excel_agrega_filass_en_memoria_un_solo_guardado():
    archivo = ArchivoSimulado()
    destino = ExcelOnlineWriter(archivo)
    campos = [Campo("numero", "numero")]
    assert destino.claves_existentes(campos) == {"outlook:demo:mensaje-1:adjunto-1"}
    destino.agregar("outlook:demo:mensaje-2:adjunto-1", "2026-10-20", ["202"])
    destino.agregar("outlook:demo:mensaje-3:adjunto-1", "2026-10-20", ["=SUM(A1:A2)"])
    destino.finalizar()
    assert len(archivo.guardados) == 1
    wb = load_workbook(BytesIO(archivo.datos))
    ws = wb["Registros"]
    assert ws.max_row == 4
    assert ws["C3"].value == "202"
    assert ws["C4"].data_type != "f"
    assert len(ws.tables) == 1
    assert list(ws.tables.values())[0].ref == "A1:C4"


def test_excel_no_sube_libro_si_no_hay_filas_nuevas():
    archivo = ArchivoSimulado()
    destino = ExcelOnlineWriter(archivo)
    destino.claves_existentes([Campo("numero", "numero")])
    destino.finalizar()
    assert archivo.guardados == []


def test_excel_rechaza_columnas_inesperadas():
    archivo = ArchivoSimulado()
    with pytest.raises(ValueError, match="no coinciden"):
        ExcelOnlineWriter(archivo).claves_existentes([Campo("otro", "otro")])


class Respuesta:
    def __init__(self, status, data):
        self.status_code = status
        self._data = data
        self.content = data if isinstance(data, bytes) else b""

    def json(self):
        return self._data


class HttpSimulado:
    def __init__(self):
        self.ultima_cabecera_subida = None
        self.datos = ejemplo_xlsx()
        self.etag = '"v1"'

    def get(self, url, headers=None, params=None, **kwargs):
        if "graph.microsoft.com" in url:
            return Respuesta(200, {
                "name": "registro.xlsx", "file": {}, "eTag": self.etag,
                "size": len(self.datos),
                "@microsoft.graph.downloadUrl": "https://empresa.sharepoint.com/descarga",
            })
        assert url == "https://empresa.sharepoint.com/descarga"
        assert "Authorization" not in headers
        assert kwargs["allow_redirects"] is False
        return Respuesta(200, self.datos)

    def put(self, url, headers, data, **kwargs):
        self.ultima_cabecera_subida = headers
        if headers["If-Match"] != self.etag:
            return Respuesta(412, {})
        self.datos = data
        return Respuesta(200, {"id": "archivo"})


class AuthSimulado:
    def acquire_token_for_client(self, scopes):
        return {"access_token": "TOKEN_FALSO"}


def test_graph_separa_bearer_y_descarga_firmada():
    http = HttpSimulado()
    graph = type("Graph", (), {"auth": AuthSimulado(), "http": http})()
    archivo = GraphExcelFile(graph, "drive001", "item001")
    contenido, etag = archivo.leer()
    assert contenido.startswith(b"PK")
    archivo.guardar(contenido, etag)
    assert http.ultima_cabecera_subida["If-Match"] == '"v1"'


def test_graph_impide_sobrescritura_si_cambio_etag():
    http = HttpSimulado()
    graph = type("Graph", (), {"auth": AuthSimulado(), "http": http})()
    archivo = GraphExcelFile(graph, "drive001", "item001")
    datos, etag = archivo.leer()
    http.etag = '"v2"'
    with pytest.raises(RuntimeError, match="Conflicto"):
        archivo.guardar(datos, etag)
    assert http.ultima_cabecera_subida is None


def test_graph_no_acepta_una_descarga_fuera_de_sharepoint():
    http = HttpSimulado()
    original = http.get
    def malicioso(url, **kwargs):
        if "graph.microsoft.com" in url:
            metadata = original(url, **kwargs).json()
            metadata["@microsoft.graph.downloadUrl"] = "https://evil.example/secret"
            return Respuesta(200, metadata)
        raise AssertionError("No se debe descargar desde un host ajeno")
    http.get = malicioso
    graph = type("Graph", (), {"auth": AuthSimulado(), "http": http})()
    with pytest.raises(ValueError, match="segura"):
        GraphExcelFile(graph, "drive001", "item001").leer()



def test_onedrive_empresa_busca_excel_por_propietario_y_ruta():
    from depurador.excel_online import resolver_ubicacion_onedrive

    class ClienteFalso:
        def __init__(self):
            self.url = None
        def get(self, url):
            self.url = url
            return {"id": "01AB123", "name": "Informe 2026.xlsx",
                    "file": {}, "parentReference": {"driveId": "b!UnDrive"}}

    cliente = ClienteFalso()
    resultado = resolver_ubicacion_onedrive(
        cliente, "informes@empresa.com", "Consolidado/Informe 2026.xlsx"
    )
    assert resultado == ("b!UnDrive", "01AB123")
    assert "/users/informes%40empresa.com/drive/root:/Consolidado/Informe%202026.xlsx:" in cliente.url


@pytest.mark.parametrize("ruta", [
    "../archivo.xlsx", "Carpeta/../archivo.xlsx", "Carpeta//archivo.xlsx",
    r"Carpeta\archivo.xlsx", "Carpeta/archivo.pdf", "",
])
def test_onedrive_rechaza_rutas_no_validas_antes_de_graph(ruta):
    from depurador.excel_online import resolver_ubicacion_onedrive

    class NoConsultar:
        def get(self, url):
            raise AssertionError("No debería comunicarse con Graph")
    with pytest.raises(ValueError, match="MS_ONEDRIVE_PATH"):
        resolver_ubicacion_onedrive(NoConsultar(), "informes@empresa.com", ruta)


def test_excel_rechaza_campos_de_longitud_equivocada():
    archivo = ArchivoSimulado()
    destino = ExcelOnlineWriter(archivo)
    destino.claves_existentes([Campo("numero", "numero")])
    with pytest.raises(ValueError, match="Cantidad"):
        destino.agregar("uno", "2026-10-20", ["100", "sobrante"])


def test_descarga_1drv_com_sin_enviar_token_bearer():
    http = HttpSimulado()
    get_original = http.get

    def falso_get(url, **kwargs):
        if "graph.microsoft.com" in url:
            metadata = get_original(url, **kwargs).json()
            metadata["@microsoft.graph.downloadUrl"] = "https://a.files.1drv.com/descarga"
            return Respuesta(200, metadata)
        assert url == "https://a.files.1drv.com/descarga"
        assert "Authorization" not in kwargs.get("headers", {})
        return Respuesta(200, http.datos)

    http.get = falso_get
    graph = type("Graph", (), {"auth": AuthSimulado(), "http": http})()
    datos, version = GraphExcelFile(graph, "b!drive", "item1").leer()
    assert datos.startswith(b"PK")
    assert version == '"v1"'
