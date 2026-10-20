"""Prueba integral offline del flujo Windows Server: Outlook -> PDF -> OneDrive Excel."""
import io
import json
import re
import threading

import fitz
from openpyxl import Workbook, load_workbook
from werkzeug.security import generate_password_hash

from depurador.web import create_app


def _pdf(numero):
    doc = fitz.open()
    p = doc.new_page()
    p.insert_text((40, 80), "Documento: " + numero)
    contenido = doc.tobytes()
    doc.close()
    return contenido


def _excel():
    wb = Workbook()
    ws = wb.active
    ws.title = "Registros"
    ws.append(["document_key", "processed_utc", "documento"])
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


class Respuesta:
    def __init__(self, codigo, datos):
        self.status_code = codigo
        self.datos = datos
        self.content = datos if isinstance(datos, bytes) else b""

    def json(self):
        return self.datos


class HttpFalso:
    def __init__(self, terminado):
        self.contenido = _excel()
        self.version = '"v1"'
        self.terminado = terminado
        self.guardados = 0

    def get(self, url, headers=None, **kwargs):
        if url.startswith("https://graph.microsoft.com/"):
            return Respuesta(200, {
                "id": "01ITEM", "name": "Datos.xlsx", "size": len(self.contenido),
                "eTag": self.version, "file": {},
                "@microsoft.graph.downloadUrl": "https://empresa.sharepoint.com/descarga",
            })
        assert url == "https://empresa.sharepoint.com/descarga"
        assert "Authorization" not in (headers or {})
        return Respuesta(200, self.contenido)

    def put(self, url, headers=None, data=None, **kwargs):
        assert url.endswith("/drives/b%21DRIVE/items/01ITEM/content")
        assert headers["If-Match"] == self.version
        self.contenido = data
        self.guardados += 1
        self.version = '"v2"'
        self.terminado.set()
        return Respuesta(200, {"id": "01ITEM"})


class AuthFalsa:
    def acquire_token_for_client(self, scopes):
        return {"access_token": "TOKEN_FICTICIO"}


class MicrosoftFalso:
    def __init__(self, terminado):
        self.auth = AuthFalsa()
        self.http = HttpFalso(terminado)
        self.leidos = []
        self.pdf = _pdf("ABC-2026")

    def get(self, url, params=None, raw=False):
        self.leidos.append((url, raw))
        if "/drive/root:/" in url:
            assert "Informes/Datos.xlsx" in url
            return {
                "id": "01ITEM", "name": "Datos.xlsx",
                "file": {}, "parentReference": {"driveId": "b!DRIVE"},
            }
        if url.endswith("/messages"):
            return {"value": [
                {"id": "m-1", "subject": "El reporte seleccionado",
                 "receivedDateTime": "2026-10-20T15:00:00Z",
                 "from": {"emailAddress": {"address": "emisor@empresa.com"}},
                 "hasAttachments": True},
                {"id": "m-2", "subject": "Otro reporte",
                 "receivedDateTime": "2026-10-19T15:00:00Z",
                 "from": {"emailAddress": {"address": "otro@empresa.com"}},
                 "hasAttachments": True},
            ]}
        if url.endswith("/attachments"):
            if "/m-1/" in url:
                return {"value": [{"id": "adj-1", "name": "informe.pdf", "size": len(self.pdf),
                                   "@odata.type": "#microsoft.graph.fileAttachment"}]}
            return {"value": [{"id": "adj-2", "name": "otro.pdf", "size": 12,
                               "@odata.type": "#microsoft.graph.fileAttachment"}]}
        if url.endswith("/attachments/adj-1/$value"):
            assert raw
            return self.pdf
        if url.endswith("/attachments/adj-2/$value"):
            raise AssertionError("Nunca se debe procesar otro correo")
        raise AssertionError("Solicitud inesperada: " + url)


def _token(client):
    html = client.get("/" if client.get("/").status_code == 200 else "/login").data
    return re.search(rb'name="_csrf" value="([^"]+)"', html).group(1).decode()


def test_recorrido_web_outlook_excel_en_onedrive(tmp_path):
    fin = threading.Event()
    microsoft = MicrosoftFalso(fin)
    reglas = tmp_path / "reglas.json"
    reglas.write_text(json.dumps({"campos": [
        {"nombre": "documento", "patron": r"Documento:\s*([A-Z0-9-]+)", "obligatorio": True}
    ]}), encoding="utf-8")
    app = create_app({
        "TESTING": True,
        "SECRET_KEY": "x" * 48,
        "ADMIN_USER": "analista",
        "ADMIN_PASSWORD_HASH": generate_password_hash("contrasena-segura-123"),
        "SESSION_COOKIE_SECURE": False,
        "MAIL_PROVIDER": "outlook",
        "TABLE_PROVIDER": "excel_online",
        "MAILBOX": "reportes@empresa.com",
        "MS_ONEDRIVE_OWNER": "datos@empresa.com",
        "MS_ONEDRIVE_PATH": "Informes/Datos.xlsx",
        "MS_EXCEL_SHEET": "Registros",
        "PDF_FIELDS_FILE": str(reglas),
    }, service_factory=lambda: (microsoft, microsoft))
    client = app.test_client()
    csrf = _token(client)
    client.post("/login", data={
        "_csrf": csrf, "usuario": "analista", "clave": "contrasena-segura-123",
    })
    csrf = _token(client)
    r = client.post("/buscar", data={"_csrf": csrf, "filtro": "from:emisor@empresa.com"})
    assert r.status_code == 302
    pagina = client.get("/")
    assert "El reporte seleccionado" in pagina.text
    assert "Otro reporte" in pagina.text
    assert "Excel Online" in pagina.text
    assert client.post("/procesar", data={
        "_csrf": csrf, "mensaje_id": "m-1",
    }).status_code == 302
    assert fin.wait(5), "El Excel simulado no se guardó"
    assert microsoft.http.guardados == 1
    wb = load_workbook(io.BytesIO(microsoft.http.contenido))
    ws = wb["Registros"]
    assert ws.max_row == 2
    assert ws["C2"].value == "ABC-2026"
    assert ws["A2"].value.startswith("outlook:reportes@empresa.com:m-1:")
    assert not any("/messages/m-2/attachments/adj-2/$value" in url for url, _ in microsoft.leidos)
